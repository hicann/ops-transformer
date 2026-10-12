#!/usr/bin/env python3
# Copyright (c) 2026 Huawei Technologies Co., Ltd.
# This program is free software, you can redistribute it and/or modify it under the terms and conditions of
# CANN Open Software License Agreement Version 2.0 (the "License").
# Please refer to the License for details. You may not use this file except in compliance with the License.
# THIS SOFTWARE IS PROVIDED ON AN "AS IS" BASIS, WITHOUT WARRANTIES OF ANY KIND, EITHER EXPRESS OR IMPLIED,
# INCLUDING BUT NOT LIMITED TO NON-INFRINGEMENT, MERCHANTABILITY, OR FITNESS FOR A PARTICULAR PURPOSE.
# See LICENSE in the root of the software repository for the full text of the License.

"""Convert this operator's parameter workbook to a TTK E2E/ACLNN workbook."""

import argparse
import hashlib
import math
import os
import re
import sys
import tempfile
from pathlib import Path
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.utils.exceptions import IllegalCharacterError, InvalidFileException

from case_schema import FIELDS, OPERATOR, convert_case


class ConversionCancelled(Exception):
    """The user declined to append; the destination remains untouched."""


def workbook_path(path):
    path = Path(path).expanduser().absolute()
    if path.suffix.lower() not in (".xlsx", ".xlsm"):
        raise ValueError(f"expected an .xlsx or .xlsm workbook: {path}")
    return path


def check_title(title):
    if not title or len(title) > 31 or re.search(r"[\\/*?:\[\]]", title):
        raise ValueError(f"invalid Excel sheet name: {title!r}")


def read_source(path, sheets, mode):
    formulas = load_workbook(path, read_only=True, data_only=False)
    values = None
    try:
        values = load_workbook(path, read_only=True, data_only=True)
        selected = sheets or [values.sheetnames[0]]
        if len(set(selected)) != len(selected):
            raise ValueError("input sheets must not be repeated")
        missing = set(selected) - set(values.sheetnames)
        if missing:
            raise ValueError(
                f"input sheets not found: {sorted(missing)}; available: {values.sheetnames}"
            )
        plans = []
        for title in selected:
            if title not in [sheet.title for sheet in values.worksheets]:
                raise ValueError(
                    f"{title}: input must be a worksheet, not a chart sheet"
                )
            rows = []
            header = None
            for row_index, (raw, original) in enumerate(
                zip(values[title].iter_rows(), formulas[title].iter_rows()), 1
            ):
                cells = [cell.value for cell in raw]
                for value, formula in zip(raw, original):
                    if formula.data_type == "f" and value.value is None:
                        raise ValueError(
                            f"{title}!{formula.coordinate}: formula has no cached value; recalculate in Excel first"
                        )
                if all(value is None or value == "" for value in cells):
                    continue
                if header is None:
                    while cells and cells[-1] in (None, ""):
                        cells.pop()
                    header = [
                        str(value).strip() if value is not None else ""
                        for value in cells
                    ]
                    if any(not field for field in header) or len(header) != len(
                        set(header)
                    ):
                        raise ValueError(
                            f"{title}: empty or duplicate input column names"
                        )
                    continue
                if any(value not in (None, "") for value in cells[len(header) :]):
                    raise ValueError(f"{title}: data extends beyond the header")
                record = dict(zip(header, cells))
                try:
                    rows.append(convert_case(record, mode))
                except (ValueError, TypeError, KeyError, ZeroDivisionError) as exc:
                    raise ValueError(f"{title}!row {row_index}: {exc}") from exc
            if header is None:
                raise ValueError(f"{title}: empty sheet")
            if not rows:
                raise ValueError(f"{title}: no testcase rows")
            plans.append((title, rows))
        return plans
    finally:
        formulas.close()
        if values is not None:
            values.close()


def fingerprint(path):
    if not path.exists():
        return None
    return hashlib.sha256(path.read_bytes()).digest()


def existing_title(book, title):
    return next(
        (name for name in book.sheetnames if name.casefold() == title.casefold()), None
    )


def table_header(sheet):
    for row in sheet.iter_rows():
        if all(cell.value in (None, "") for cell in row):
            continue
        fields = [cell.value for cell in row]
        while fields and fields[-1] in (None, ""):
            fields.pop()
        if any(not isinstance(field, str) or not field.strip() for field in fields):
            raise ValueError(f"{sheet.title}: invalid destination header")
        fields = [field.strip() for field in fields]
        if len(fields) != len(set(fields)):
            raise ValueError(f"{sheet.title}: duplicate destination columns")
        return row[0].row, fields
    return 1, []


def prepare_append(book, plans, fields, explicit_sheet):
    prepared = []
    for title, rows in plans:
        check_title(title)
        name = existing_title(book, title)
        if name is None:
            if len(rows) + 1 > 1048576:
                raise ValueError(f"{title}: cases exceed the Excel row limit")
            prepared.append((title, rows, 1, list(fields), False))
            continue
        sheet = book[name]
        if name not in [sheet.title for sheet in book.worksheets]:
            raise ValueError(
                f"{name}: destination must be a worksheet, not a chart sheet"
            )
        if sheet.protection.sheet:
            raise ValueError(f"{name}: destination sheet is protected")
        if sheet.tables:
            raise ValueError(
                f"{name}: Excel structured tables are not supported; use a plain testcase sheet"
            )
        header_row, header = table_header(sheet)
        if header:
            required = {
                "testcase_name",
                "api_name",
                "tensor_view_shapes",
                "tensor_dtypes",
                "attributes",
            }
            if not required.issubset(header):
                raise ValueError(f"{name}: existing sheet is not a TTK testcase table")
            if ("output_tensor_indexes" in header) != (
                "output_tensor_indexes" in fields
            ):
                raise ValueError(f"{name}: cannot mix E2E and ACLNN rows in one sheet")
        header += [field for field in fields if field not in header]
        if sheet.merged_cells.ranges:
            raise ValueError(
                f"{name}: merged cells are not supported in a testcase table"
            )
        last_row = max(
            (
                cell.row
                for row in sheet.iter_rows()
                for cell in row
                if cell.value is not None
            ),
            default=0,
        )
        if max(last_row, header_row) + len(rows) > 1048576:
            raise ValueError(f"{name}: appended cases exceed the Excel row limit")
        prepared.append((name, rows, header_row, header, not explicit_sheet))
    for title, rows, _, header, _ in prepared:
        if len(header) > 16384:
            raise ValueError(f"{title}: too many Excel columns")
        for row in rows:
            for field, value in row.items():
                if not isinstance(value, (str, int, float, bool)) and value is not None:
                    raise ValueError(f"{title}: {field} is not an Excel scalar")
                if isinstance(value, float) and not math.isfinite(value):
                    raise ValueError(
                        f"{title}: {field} requires literal text for inf/nan"
                    )
                if isinstance(value, str) and len(value) > 32767:
                    raise ValueError(
                        f"{title}: {field} exceeds Excel's 32767-character cell limit"
                    )
    return prepared


def write_text(cell, value):
    cell.value = value
    # TTK fields are literal text, including names that begin with '='.
    if isinstance(value, str):
        cell.data_type = "s"


def publish(destination, plans, fields, explicit_sheet=False, source=None):
    if destination.is_symlink():
        raise ValueError(f"destination must not be a symbolic link: {destination}")
    if destination.exists() and not destination.is_file():
        raise ValueError(f"destination is not a file: {destination}")
    before = fingerprint(destination)
    if before is None:
        book = Workbook()
        book.remove(book.active)
        if destination.suffix.lower() == ".xlsm":
            if source is None or source.suffix.lower() != ".xlsm":
                raise ValueError(
                    "creating .xlsm output requires an .xlsm source; otherwise use .xlsx"
                )
            macro_source = load_workbook(source, keep_vba=True)
            try:
                book.vba_archive = macro_source.vba_archive
                book.code_name = macro_source.code_name
            finally:
                macro_source.close()
    else:
        book = load_workbook(
            destination, keep_vba=destination.suffix.lower() == ".xlsm", rich_text=True
        )
    temporary = None
    try:
        prepared = prepare_append(book, plans, fields, explicit_sheet)
        for title, rows, _, _, ask in prepared:
            if ask:
                try:
                    answer = input(
                        f"Append {len(rows)} cases to existing sheet {title!r} in {destination.name!r}? [y/N]: "
                    )
                except EOFError as exc:
                    raise ConversionCancelled(
                        "no confirmation received; no changes written"
                    ) from exc
                if answer.strip() not in ("yes", "y", "Y"):
                    raise ConversionCancelled("append declined; no changes written")
        for title, rows, header_row, header, _ in prepared:
            sheet = (
                book[title] if title in book.sheetnames else book.create_sheet(title)
            )
            for column, field in enumerate(header, 1):
                write_text(sheet.cell(header_row, column), field)
            last = max(
                (
                    cell.row
                    for row in sheet.iter_rows()
                    for cell in row
                    if cell.value is not None
                ),
                default=header_row,
            )
            for index, record in enumerate(rows, last + 1):
                for column, field in enumerate(header, 1):
                    write_text(sheet.cell(index, column), record.get(field, ""))
        destination.parent.mkdir(parents=True, exist_ok=True)
        with tempfile.NamedTemporaryFile(
            dir=destination.parent,
            prefix=f".{destination.stem}.",
            suffix=destination.suffix,
            delete=False,
        ) as stream:
            temporary = Path(stream.name)
        book.save(temporary)
        if destination.is_symlink() or fingerprint(destination) != before:
            raise ValueError(
                "destination changed during conversion; no changes written"
            )
        if before is not None:
            os.chmod(temporary, destination.stat().st_mode & 0o777)
        os.replace(temporary, destination)
        temporary = None
    finally:
        book.close()
        if temporary is not None:
            temporary.unlink(missing_ok=True)


def build_parser():
    parser = argparse.ArgumentParser(
        description=f"Convert {OPERATOR} Excel parameters to TTK Excel cases."
    )
    parser.add_argument(
        "-i", "--input", required=True, type=Path, help="source .xlsx/.xlsm workbook"
    )
    parser.add_argument(
        "-s",
        "--sheet",
        action="append",
        nargs="+",
        help="source sheet names; repeatable; default: first sheet",
    )
    parser.add_argument(
        "--mode",
        choices=("e2e", "aclnn"),
        default="e2e",
        help="TTK table format (default: e2e)",
    )
    output = parser.add_mutually_exclusive_group(required=True)
    output.add_argument(
        "--output-dir",
        type=Path,
        help="directory; retain the input workbook filename and sheet names",
    )
    output.add_argument(
        "-o",
        "--output",
        type=Path,
        help="destination workbook; retain source sheet names and append",
    )
    parser.add_argument(
        "--output-sheet",
        help="with --output, append all selected source sheets to this target sheet",
    )
    return parser


def main(argv=None):
    parser = build_parser()
    args = parser.parse_args(argv)
    try:
        source = workbook_path(args.input)
        if not source.is_file():
            raise ValueError(f"source workbook not found: {source}")
        if args.output_sheet is not None and args.output is None:
            raise ValueError("--output-sheet requires --output")
        destination = workbook_path(
            args.output
            if args.output is not None
            else args.output_dir.expanduser() / source.name
        )
        if (
            args.output_dir is not None
            and args.output_dir.exists()
            and not args.output_dir.is_dir()
        ):
            raise ValueError(f"output directory is not a directory: {args.output_dir}")
        if source.resolve() == destination.resolve() or (
            destination.exists() and os.path.samefile(source, destination)
        ):
            raise ValueError("input and output workbooks must be different")
        if args.output_sheet is not None:
            check_title(args.output_sheet)
        sheets = [name for group in args.sheet or [] for name in group]
        plans = read_source(source, sheets, args.mode)
        if args.output_sheet is not None:
            plans = [(args.output_sheet, [row for _, rows in plans for row in rows])]
        fields = list(
            FIELDS + (("output_tensor_indexes",) if args.mode == "aclnn" else ())
        )
        for _, rows in plans:
            for row in rows:
                fields.extend(field for field in row if field not in fields)
        publish(
            destination,
            plans,
            fields,
            explicit_sheet=args.output_sheet is not None,
            source=source,
        )
        print(
            f"{OPERATOR}: wrote {sum(len(rows) for _, rows in plans)} {args.mode} cases to {destination}"
        )
        for title, rows in plans:
            print(f"  {title}: {len(rows)} cases")
        return 0
    except (
        OSError,
        ValueError,
        KeyError,
        BadZipFile,
        ParseError,
        InvalidFileException,
        IllegalCharacterError,
        ConversionCancelled,
    ) as exc:
        print(f"error: {exc}", file=sys.stderr)
        return 1
    except KeyboardInterrupt:
        print("error: conversion interrupted", file=sys.stderr)
        return 1


if __name__ == "__main__":
    sys.exit(main())
