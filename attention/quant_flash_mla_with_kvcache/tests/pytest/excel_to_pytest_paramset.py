#!/usr/bin/python3
# -*- coding: utf-8 -*-
# -----------------------------------------------------------------------------------------------------------
# Copyright (c) 2026 Huawei Technologies Co., Ltd.
# This program is free software, you can redistribute it and/or modify it under the terms and conditions of
# CANN Open Software License Agreement Version 2.0 (the "License").
# Please refer to the License for details. You may not use this file except in compliance with the License.
# THIS SOFTWARE IS PROVIDED ON AN "AS IS" BASIS, WITHOUT WARRANTIES OF ANY KIND, EITHER EXPRESS OR IMPLIED,
# INCLUDING BUT NOT LIMITED TO NON-INFRINGEMENT, MERCHANTABILITY, OR FITNESS FOR A PARTICULAR PURPOSE.
# See LICENSE in the root of the software repository for the full text of the License.
# -----------------------------------------------------------------------------------------------------------
"""把 xlsx 参数表转换为 quant_flash_mla_paramset_debug.py 需要的 TEST_PARAMS 格式。

环境：无 openpyxl/pandas，xlsx 解析用内置极简读取器（zipfile + minidom）。

用法：
  python3 excel_to_pytest_paramset.py --xlsx red_modified.xlsx
  python3 excel_to_pytest_paramset.py --xlsx xx.xlsx --sheet 1
  python3 excel_to_pytest_paramset.py --xlsx xx.xlsx --dry-run

字段映射（Excel colon -> paramset key）：
  B                 -> B
  Q_N / num_heads_q -> N_q
  KV_N/num_heads_kv -> N_kv
  Q_D / head_dim    -> D
  head_dim_v        -> D_V
  head_dim_rope / D_rope -> D_rope
  cu_seqlens_q_value-> cu_seqlens_q
  cache_seqlens_value-> cache_seqlens
  seqused_q_value   -> seqused_q
  max_seqlen_q      -> max_seqlen_q
  max_seqlen_kv     -> max_seqlen_kv
  layout_kv 前缀 PA_ -> enable_pa (True/False)
  layout_kv         -> kv_cache_layout
  block_size        -> block_size
  mask_mode         -> mask_mode
  p_scale_value     -> p_scale (空则默认 1.0)
  softmax_scale     -> softmax_scale
  return_softmax_lse-> enable_lse
  quant_mode        -> quant_mode
"""

import argparse
import os
import re
import zipfile
import xml.dom.minidom as minidom

_HERE = os.path.dirname(os.path.abspath(__file__))
_DEFAULT_SHEET = 1
_DEFAULT_OUT = os.path.join(_HERE, "quant_flash_mla_paramset_debug.py")


# ---------------------------------------------------------------------------
# 内置极简 xlsx 读取器（zipfile + minidom，无 openpyxl 依赖）
# ---------------------------------------------------------------------------
_CELL_RE = re.compile(r"^([A-Z]+)(\d+)$")


def col_to_idx(col):
    n = 0
    for ch in col:
        n = n * 26 + (ord(ch.upper()) - ord("A") + 1)
    return n - 1


def _parse_ref(ref):
    m = _CELL_RE.match(ref)
    if not m:
        raise ValueError(f"bad cell ref: {ref!r}")
    return col_to_idx(m.group(1)), int(m.group(2)) - 1


def _cell_text(cell, shared_strings):
    t = cell.getAttribute("t")
    v_els = cell.getElementsByTagName("v")
    if t == "s" and v_els:
        idx = int(v_els[0].firstChild.data)
        return shared_strings[idx]
    if t == "b" and v_els:
        return "True" if v_els[0].firstChild.data == "1" else "False"
    if t == "inlineStr":
        is_els = cell.getElementsByTagName("is")
        if not is_els:
            return None
        ts = is_els[0].getElementsByTagName("t")
        return "".join(tt.firstChild.data if tt.firstChild else "" for tt in ts)
    if v_els:
        return v_els[0].firstChild.data
    return None


def _resolve_sheet_xml(z, sheet):
    if isinstance(sheet, int):
        return f"xl/worksheets/sheet{sheet}.xml"
    wb_xml = z.read("xl/workbook.xml").decode("utf-8")
    wb_doc = minidom.parseString(wb_xml)
    rid = None
    for s in wb_doc.getElementsByTagName("sheet"):
        if s.getAttribute("name") == sheet:
            rid = s.getAttribute("r:id") or s.getAttribute(
                "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
            )
            break
    if rid is None:
        names = [s.getAttribute("name") for s in wb_doc.getElementsByTagName("sheet")]
        raise ValueError(f"sheet name {sheet!r} not found. available: {names}")
    rels_xml = z.read("xl/_rels/workbook.xml.rels").decode("utf-8")
    rels_doc = minidom.parseString(rels_xml)
    for rel in rels_doc.getElementsByTagName("Relationship"):
        if rel.getAttribute("Id") == rid:
            target = rel.getAttribute("Target")
            return target.lstrip("/") if target.startswith("/") else "xl/" + target
    raise ValueError(f"r:id {rid!r} not found in workbook.xml.rels")


def _parse_sheet_xml(sheet_xml, shared_strings):
    sh = minidom.parseString(sheet_xml)
    rows = sh.getElementsByTagName("row")
    rows = sorted(rows, key=lambda r: int(r.getAttribute("r")))
    parsed_rows = []
    for row in rows:
        cells = {}
        for c in row.getElementsByTagName("c"):
            col_idx, _ = _parse_ref(c.getAttribute("r"))
            cells[col_idx] = _cell_text(c, shared_strings)
        parsed_rows.append(cells)
    return parsed_rows


def load_sheet(xlsx_path, sheet=1):
    try:
        with zipfile.ZipFile(xlsx_path) as z:
            ss_xml = z.read("xl/sharedStrings.xml").decode("utf-8")
            sheet_path = _resolve_sheet_xml(z, sheet)
            sheet_xml = z.read(sheet_path).decode("utf-8")
    except (KeyError, FileNotFoundError):
        import openpyxl

        wb = openpyxl.load_workbook(xlsx_path, data_only=True)
        if isinstance(sheet, int):
            ws = (
                wb.worksheets[sheet - 1]
                if 1 <= sheet <= len(wb.worksheets)
                else wb.active
            )
        else:
            ws = wb[sheet]
        parsed_rows = []
        for row in ws.iter_rows():
            cells = {}
            for cell in row:
                v = cell.value
                if v is None:
                    continue
                cells[cell.col_idx - 1] = str(v) if not isinstance(v, str) else v
            parsed_rows.append(cells)
        return parsed_rows

    ss_doc = minidom.parseString(ss_xml)
    shared_strings = []
    for si in ss_doc.getElementsByTagName("si"):
        ts = si.getElementsByTagName("t")
        shared_strings.append(
            "".join(t.firstChild.data if t.firstChild else "" for t in ts)
        )
    return _parse_sheet_xml(sheet_xml, shared_strings)


def list_sheet_names(xlsx_path):
    try:
        with zipfile.ZipFile(xlsx_path) as z:
            wb_xml = z.read("xl/workbook.xml").decode("utf-8")
    except (KeyError, FileNotFoundError):
        import openpyxl

        wb = openpyxl.load_workbook(xlsx_path, read_only=True)
        return wb.sheetnames
    wb_doc = minidom.parseString(wb_xml)
    return [s.getAttribute("name") for s in wb_doc.getElementsByTagName("sheet")]


_FILE_HEADER = """\
#!/usr/bin/python3
# -*- coding: utf-8 -*-
# -----------------------------------------------------------------------------------------------------------
# Copyright (c) 2026 Huawei Technologies Co., Ltd.
# This program is free software, you can redistribute it and/or modify it under the terms and conditions of
# CANN Open Software License Agreement Version 2.0 (the "License").
# Please refer to the License for details. You may not use this file except in compliance with the License.
# THIS SOFTWARE IS PROVIDED ON AN "AS IS" BASIS, WITHOUT WARRANTIES OF ANY KIND, EITHER EXPRESS OR IMPLIED,
# INCLUDING BUT NOT LIMITED TO NON-INFRINGEMENT, MERCHANTABILITY, OR FITNESS FOR A PARTICULAR PURPOSE.
# See LICENSE in the root of the software repository for the full text of the License.
# -----------------------------------------------------------------------------------------------------------
# Auto-generated by excel_to_pytest_paramset.py — do not edit manually.

from quant_flash_mla_paramset_common import expand_paramset_to_cases

"""

_FILE_FOOTER = """
CASES = expand_paramset_to_cases(TEST_PARAMS)
"""


class _Columns:
    def __init__(self, header_row):
        self._by_name = {}
        for col_idx, text in sorted(header_row.items()):
            if text is None:
                continue
            key = str(text).strip()
            if key and key not in self._by_name:
                self._by_name[key] = col_idx

    def get(self, name):
        return self._by_name.get(name)


def _strip(s):
    if s is None:
        return None
    s = s.strip()
    return s or None


def _to_int(s):
    s = _strip(s)
    if s is None:
        return None
    return int(float(s))


def _to_float(s):
    s = _strip(s)
    if s is None:
        return None
    return float(s)


def _to_int_list(s):
    s = _strip(s)
    if s is None:
        return None
    parts = [p.strip() for p in s.split(",") if p.strip() != ""]
    return [int(float(p)) for p in parts]


def _to_bool(s):
    s = _strip(s)
    if s is None:
        return None
    return s in ("True", "TRUE", "1", "true")


def _parse_row(row, cols):
    params = {}
    params["B"] = _to_int(row.get(cols.get("B")))
    params["N_q"] = _to_int(
        row.get(cols.get("Q_N")) or row.get(cols.get("num_heads_q"))
    )
    params["N_kv"] = _to_int(
        row.get(cols.get("KV_N")) or row.get(cols.get("num_heads_kv"))
    )
    params["D"] = _to_int(row.get(cols.get("Q_D")) or row.get(cols.get("head_dim")))
    params["D_V"] = _to_int(row.get(cols.get("head_dim_v")))
    params["D_rope"] = _to_int(
        row.get(cols.get("head_dim_rope")) or row.get(cols.get("D_rope"))
    )

    # Excel Q_D 是 nope+rope 合并值 (576)，golden 需要 nope(512)+rope(64)
    # 如果 D_rope 未填且 D > 512，推断 D_rope = D - 512, D = 512
    if params["D_rope"] is None and params["D"] and params["D"] > 512:
        params["D_rope"] = params["D"] - 512
        params["D"] = 512

    params["cu_seqlens_q"] = _to_int_list(row.get(cols.get("cu_seqlens_q_value")))
    params["cache_seqlens"] = _to_int_list(row.get(cols.get("cache_seqlens_value")))
    params["seqused_q"] = _to_int_list(row.get(cols.get("seqused_q_value")))

    params["max_seqlen_q"] = _to_int(row.get(cols.get("max_seqlen_q")))
    params["max_seqlen_kv"] = _to_int(row.get(cols.get("max_seqlen_kv")))

    layout_kv = _strip(row.get(cols.get("layout_kv")))
    params["enable_pa"] = bool(layout_kv and layout_kv.startswith("PA_"))
    params["kv_cache_layout"] = layout_kv
    params["block_size"] = _to_int(
        row.get(cols.get("block_size")) or row.get(cols.get("blocksize"))
    )
    params["mask_mode"] = _to_int(row.get(cols.get("mask_mode")))

    p_scale = _to_float(row.get(cols.get("p_scale_value")))
    params["p_scale"] = 1.0 if p_scale is None else p_scale

    params["softmax_scale"] = _to_float(row.get(cols.get("softmax_scale")))
    params["enable_lse"] = bool(
        _to_bool(row.get(cols.get("return_softmax_lse"))) or False
    )
    params["quant_mode"] = _to_int(row.get(cols.get("quant_mode")))
    return params


_LIST_FIELDS = {"cu_seqlens_q", "cache_seqlens", "seqused_q"}


def _render_value(key, value):
    if key in _LIST_FIELDS:
        if value is None:
            return "[None]"
        return "[" + "[" + ", ".join(str(v) for v in value) + "]" + "]"
    if isinstance(value, bool):
        return "[True]" if value else "[False]"
    if isinstance(value, str):
        return f'["{value}"]'
    if isinstance(value, float):
        if value == int(value):
            return f"[{value:.1f}]"
        return f"[{value!r}]"
    if value is None:
        return "[None]"
    return f"[{value}]"


_FIELDS = [
    "B",
    "N_q",
    "N_kv",
    "D",
    "D_V",
    "D_rope",
    "cu_seqlens_q",
    "cache_seqlens",
    "seqused_q",
    "max_seqlen_q",
    "max_seqlen_kv",
    "enable_pa",
    "kv_cache_layout",
    "block_size",
    "mask_mode",
    "p_scale",
    "softmax_scale",
    "enable_lse",
    "quant_mode",
]


def build_test_params(data_rows, cols):
    test_params = {}
    for row in data_rows:
        name = _strip(row.get(cols.get("testcase_name")))
        if not name:
            continue
        parsed = _parse_row(row, cols)
        test_params[name] = {k: _render_value(k, parsed[k]) for k in _FIELDS}
    return test_params


def render_file(test_params):
    lines = [_FILE_HEADER, "TEST_PARAMS = {\n"]
    for name, params in test_params.items():
        lines.append(f'    "{name}": {{\n')
        for key, val in params.items():
            lines.append(f'        "{key}": {val},\n')
        lines.append("    },\n")
    lines.append("}\n")
    lines.append(_FILE_FOOTER)
    return "".join(lines)


def main():
    parser = argparse.ArgumentParser(
        description="xlsx -> quant_flash_mla_paramset_debug.py 格式"
    )
    parser.add_argument("--xlsx", required=True, help="输入 xlsx 路径")
    parser.add_argument(
        "--sheet", default=_DEFAULT_SHEET, help=f"sheet 名/序号 (默认 {_DEFAULT_SHEET})"
    )
    parser.add_argument("--output", default=_DEFAULT_OUT, help="输出 paramset 文件路径")
    parser.add_argument("--dry-run", action="store_true", help="只打印, 不写文件")
    parser.add_argument("--list-sheets", action="store_true", help="列出所有 sheet 名")
    args = parser.parse_args()

    if args.list_sheets:
        for i, nm in enumerate(list_sheet_names(args.xlsx), 1):
            print(f"  [{i}] {nm}")
        return

    sheet = int(args.sheet) if str(args.sheet).strip().isdigit() else args.sheet
    print(f"[excel_to_pytest_paramset] reading sheet={sheet!r} from {args.xlsx}")
    rows = load_sheet(args.xlsx, sheet)
    if not rows:
        raise RuntimeError(f"xlsx sheet {sheet!r} has no rows")
    header_row, data_rows = rows[0], rows[1:]
    cols = _Columns(header_row)
    if cols.get("testcase_name") is None:
        raise RuntimeError("sheet has no testcase_name column")

    test_params = build_test_params(data_rows, cols)
    if not test_params:
        raise RuntimeError(f"sheet {sheet!r} has no valid case rows")

    content = render_file(test_params)
    print(f"  generated {len(test_params)} case(s)")
    if args.dry_run:
        print(content)
        return

    out_path = os.path.abspath(args.output)
    with open(out_path, "w", encoding="utf-8") as f:
        f.write(content)
    print(f"wrote {out_path}")


if __name__ == "__main__":
    main()
