# Copyright (c) 2026 Huawei Technologies Co., Ltd.
# This program is free software, you can redistribute it and/or modify it under the terms and conditions of
# CANN Open Software License Agreement Version 2.0 (the "License").
# Please refer to the License for details. You may not use this file except in compliance with the License.
# THIS SOFTWARE IS PROVIDED ON AN "AS IS" BASIS, WITHOUT WARRANTIES OF ANY KIND, EITHER EXPRESS OR IMPLIED,
# INCLUDING BUT NOT LIMITED TO NON-INFRINGEMENT, MERCHANTABILITY, OR FITNESS FOR A PARTICULAR PURPOSE.
# See LICENSE in the root of the software repository for the full text of the License.

import os
import hashlib
import json
from pathlib import Path

import pytest
import torch

from qsli_test_utils import (
    QsliCaseSelector,
    compare_qsli_all_rows,
    compare_qsli_case,
    convert_qsli_case_to_tnd,
    generate_qsli_test_data,
    run_qsli_case,
)

PT_DIR = os.environ.get("QSLI_TESTCASE_DIR", "pt_path")
PT_LIST = os.environ.get("QSLI_PT_FILE_LIST", "")
CASE_NAMES = os.environ.get("QSLI_CASE_NAMES", "")
CASE_INDEXES = os.environ.get("QSLI_CASE_INDEXES", "")
SINGLE_PATH = os.environ.get("QSLI_TESTCASE_PATH", "")

try:
    TESTCASE_FILES = QsliCaseSelector.resolve(
        PT_DIR,
        explicit_files=SINGLE_PATH or PT_LIST,
        case_names=CASE_NAMES,
        case_indexes=CASE_INDEXES,
    )
except ValueError:
    if any(
        (
            SINGLE_PATH,
            PT_LIST,
            CASE_NAMES,
            CASE_INDEXES,
            os.environ.get("QSLI_TESTCASE_DIR"),
        )
    ):
        raise
    TESTCASE_FILES = []


@pytest.mark.ci
@pytest.mark.npu
@pytest.mark.parametrize("testcase_file", TESTCASE_FILES, ids=lambda x: Path(x).stem)
def test_qsli_batch(testcase_file):
    print(f"执行用例：{testcase_file}")
    case_data = torch.load(
        testcase_file, map_location="cpu", weights_only=False, mmap=True
    )
    actual = run_qsli_case(case_data, int(os.environ.get("QSLI_DEVICE_ID", "0")))
    compare = (
        compare_qsli_all_rows if case_data["cpu_result"] is None else compare_qsli_case
    )
    result, fulfill = compare(case_data, actual)
    print(f"result: {result}")
    print(f"fulfill_percent: {fulfill}")
    print("result_return_value: Pass")


SELECTION_PATH = os.environ.get("QSLI_EXCEL_SELECTION", "")
EXCEL_CASES = (
    json.loads(Path(SELECTION_PATH).read_text())["cases"] if SELECTION_PATH else []
)


@pytest.mark.ci
@pytest.mark.npu
@pytest.mark.parametrize("selection", EXCEL_CASES, ids=lambda value: value["case_name"])
def test_qsli_excel_layouts(selection):
    from openpyxl import load_workbook
    from batch.quant_sparse_lightning_indexer_pt_save import _value

    workbook_path = Path(__file__).parent / "excel" / selection["workbook"]
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        sheet = workbook[selection["sheet"]]
        header = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
        values = next(
            sheet.iter_rows(
                min_row=selection["row"], max_row=selection["row"], values_only=True
            )
        )
        source = dict(zip(header, values))
    finally:
        workbook.close()
    assert source["Testcase_Name"] == selection["case_name"]
    params = {
        name: _value(value) for name, value in source.items() if name != "Testcase_Name"
    }
    pa_case = generate_qsli_test_data(params, defer_golden=True)
    tnd_case = convert_qsli_case_to_tnd(pa_case)
    device = int(os.environ.get("QSLI_DEVICE_ID", "0"))
    pa_actual = run_qsli_case(pa_case, device)
    tnd_actual = run_qsli_case(tnd_case, device)
    result, fulfill = compare_qsli_all_rows(
        pa_case, pa_actual, peer_case=tnd_case, peer_actual=tnd_actual
    )
    fingerprints = {}
    for layout, case in (("PA_BBND", pa_case), ("TND", tnd_case)):
        fingerprints[layout] = {
            name: dict(
                shape=list(tensor.shape),
                dtype=str(tensor.dtype),
                stride=list(tensor.stride()),
                sha256=hashlib.sha256(
                    tensor.contiguous().numpy().tobytes()
                ).hexdigest(),
            )
            for name, tensor in case["tensors"].items()
            if isinstance(tensor, torch.Tensor) and not name.startswith("_")
        }
    summary = dict(
        source=selection,
        params=pa_case["params"],
        result=result,
        layouts=["PA_BBND", "TND"],
        rows_checked=pa_case["tensors"]["q"].shape[0],
        golden_chunk_rows=32,
        fulfill_min=fulfill,
        input_fingerprints=fingerprints,
        outputs_identical={
            name: torch.equal(pa_actual[name], tnd_actual[name]) for name in pa_actual
        },
        workbook_sha256=hashlib.sha256(workbook_path.read_bytes()).hexdigest(),
    )
    output = os.environ.get("QSLI_VALIDATION_OUTPUT")
    if output:
        directory = Path(output)
        directory.mkdir(parents=True, exist_ok=True)
        (directory / (selection["case_name"] + ".json")).write_text(
            json.dumps(summary, ensure_ascii=False, indent=2) + "\n"
        )
    print(
        "QSLI_PAIRED_RESULT",
        json.dumps(
            {
                name: value
                for name, value in summary.items()
                if name != "input_fingerprints"
            }
        ),
    )
