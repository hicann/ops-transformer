# Copyright (c) 2026 Huawei Technologies Co., Ltd.
# This program is free software, you can redistribute it and/or modify it under the terms and conditions of
# CANN Open Software License Agreement Version 2.0 (the "License").
# Please refer to the License for details. You may not use this file except in compliance with the License.
# THIS SOFTWARE IS PROVIDED ON AN "AS IS" BASIS, WITHOUT WARRANTIES OF ANY KIND, EITHER EXPRESS OR IMPLIED,
# INCLUDING BUT NOT LIMITED TO NON-INFRINGEMENT, MERCHANTABILITY, OR FITNESS FOR A PARTICULAR PURPOSE.
# See LICENSE in the root of the software repository for the full text of the License.

import ast
import math
from pathlib import Path

from test_case import DEFAULTS, MODES

BOOL_FIELDS = {
    "return_softmax_lse",
    "kv_axis0_noncontiguous",
    "omit_ori_topk_length",
    "omit_cmp_topk_length",
}
INT_FIELDS = {
    "B",
    "K1",
    "K2",
    "block_size1",
    "block_size2",
    "quant_mode",
    "seed",
    "N1",
    "N2",
    "D",
    "rope_head_dim",
    "cmp_ratio",
}
LENGTH_FIELDS = {"S1", "S2", "S2C"}
QUERY_FIELDS = {"cu_seqlens_q", "seqused_q"}
TOPK_FIELDS = {"ori_topk_length", "cmp_topk_length"}
FIXED_FIELDS = {
    "layout_q": "TND",
    "layout_kv": "PA_BBND",
    "N1": 64,
    "N2": 1,
    "D": 512,
    "rope_head_dim": 64,
    "quant_mode": 1,
}
ALIASES = {"K": "K2"}
COMMENT_FIELDS = {"备注"}
SUPPORTED_FIELDS = (
    set(DEFAULTS)
    | BOOL_FIELDS
    | INT_FIELDS
    | LENGTH_FIELDS
    | set(FIXED_FIELDS)
    | set(ALIASES)
    | COMMENT_FIELDS
    | {"Testcase_Name", "template_run_mode", "softmax_scale"}
)


def _empty(value):
    return value is None or (
        isinstance(value, str) and value.strip() in ("", "None", "null")
    )


def _integer(value):
    if isinstance(value, bool):
        raise ValueError("expected an integer, got a boolean")
    number = float(value)
    if not math.isfinite(number) or not number.is_integer():
        raise ValueError(f"expected an integer, got {value!r}")
    return int(number)


def _boolean(value):
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)) and value in (0, 1):
        return bool(value)
    if isinstance(value, str) and value.lower() in ("true", "false", "0", "1"):
        return value.lower() in ("true", "1")
    raise ValueError(f"expected TRUE/FALSE or 1/0, got {value!r}")


def _normalize(values):
    p = {}
    for original, value in values.items():
        if original in COMMENT_FIELDS or _empty(value):
            continue
        if original not in SUPPORTED_FIELDS:
            raise ValueError(
                f"unsupported column {original!r}; see the DSL Excel template"
            )
        field = ALIASES.get(original, original)
        if isinstance(value, str):
            value = value.strip()
        try:
            if field in BOOL_FIELDS:
                value = _boolean(value)
            elif field in INT_FIELDS:
                value = _integer(value)
            elif field in LENGTH_FIELDS:
                if isinstance(value, str) and value.startswith("["):
                    value = ast.literal_eval(value)
                value = (
                    [_integer(x) for x in value]
                    if isinstance(value, list)
                    else _integer(value)
                )
            elif field in QUERY_FIELDS:
                if isinstance(value, str):
                    value = ast.literal_eval(value)
                if not isinstance(value, list):
                    raise ValueError("expected a one-dimensional integer list")
            elif field in TOPK_FIELDS:
                if isinstance(value, str) and value.startswith("["):
                    value = ast.literal_eval(value)
                if not isinstance(value, list):
                    value = _integer(value)
            elif field == "softmax_scale":
                value = float(value)
                if not math.isfinite(value):
                    raise ValueError("must be finite")
        except (TypeError, ValueError, SyntaxError) as exc:
            raise ValueError(f"{original}: {exc}") from exc
        if field in p and p[field] != value:
            raise ValueError(f"conflicting values for {original} and {field}")
        p[field] = value
    if not isinstance(p.get("Testcase_Name"), str):
        raise ValueError("Testcase_Name must be a non-empty string")
    if p.get("template_run_mode") not in MODES:
        raise ValueError(f"template_run_mode must be one of {MODES}")
    for field, expected in FIXED_FIELDS.items():
        if field in p and p[field] != expected:
            raise ValueError(f"{field} supports only {expected!r}")
    merged = {**DEFAULTS, **p}
    for field in ("B", "K1", "K2", "block_size1", "block_size2"):
        if merged[field] <= 0:
            raise ValueError(f"{field} must be positive")
    for field in LENGTH_FIELDS:
        if field == "S1" and any(merged[name] is not None for name in QUERY_FIELDS):
            continue
        value = merged[field]
        values = value if isinstance(value, list) else [value]
        if (
            not values
            or (isinstance(value, list) and len(value) != merged["B"])
            or any(x <= 0 for x in values)
        ):
            raise ValueError(
                f"{field} must be positive: a scalar or a list of B lengths"
            )
    if "cmp_ratio" in p:
        if p["cmp_ratio"] <= 0:
            raise ValueError("cmp_ratio must be positive")
        s2 = merged["S2"]
        derived = (
            [x // p["cmp_ratio"] for x in s2]
            if isinstance(s2, list)
            else s2 // p["cmp_ratio"]
        )
        if "S2C" in p and p["S2C"] != derived:
            raise ValueError("S2C conflicts with S2 // cmp_ratio")
        if min(derived if isinstance(derived, list) else [derived]) <= 0:
            raise ValueError("S2 // cmp_ratio must be positive")
        p["S2C"] = derived
    from mqsmla_golden import _topk_lengths, query_lengths

    _, cumulative, _ = query_lengths(
        merged["B"], merged["S1"], merged["cu_seqlens_q"], merged["seqused_q"]
    )
    t1 = int(cumulative[-1])
    for field, width in (("ori_topk_length", "K1"), ("cmp_topk_length", "K2")):
        if p.get(field) is not None:
            try:
                _topk_lengths(t1, merged[width], p[field])
            except (ValueError, TypeError, RuntimeError) as exc:
                raise ValueError(f"{field}: {exc}") from exc
    for field, length in (
        ("ori_kv_topk_mode", "ori_topk_length"),
        ("cmp_kv_topk_mode", "cmp_topk_length"),
    ):
        if merged[length] is None and merged[field] not in ("fullK", "random"):
            raise ValueError(f"{field} supports only fullK/random")
    if merged["dist"] != "norm":
        raise ValueError("dist supports only 'norm'")
    if p.get("seed", DEFAULTS["seed"]) < 0:
        raise ValueError("seed must be non-negative")
    # Fixed contract columns have been checked; the runner already uses them.
    for field in FIXED_FIELDS.keys() - {"quant_mode"}:
        p.pop(field, None)
    p.pop("cmp_ratio", None)
    return p


def load_excel_test_cases(path, sheet="decode"):
    """Return normalized params; errors include workbook, sheet and Excel row."""
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ValueError(
            "Excel input requires openpyxl: python -m pip install openpyxl"
        ) from exc
    path = Path(path).expanduser().resolve()
    if not path.is_file():
        raise ValueError(f"Excel file not found: {path}")
    try:
        workbook = load_workbook(path, read_only=True, data_only=False)
    except Exception as exc:
        raise ValueError(f"Cannot read Excel file {path}: {exc}") from exc
    try:
        if sheet not in workbook.sheetnames:
            raise ValueError(
                f"{path}: sheet {sheet!r} not found; available: {workbook.sheetnames}"
            )
        rows = workbook[sheet].iter_rows()
        headers = [
            str(cell.value).strip() if cell.value is not None else ""
            for cell in next(rows, ())
        ]
        named = [name for name in headers if name]
        if len(named) != len(set(named)):
            raise ValueError(f"{path} [{sheet}]: duplicate column names")
        if not {"Testcase_Name", "template_run_mode"} <= set(named):
            raise ValueError(
                f"{path} [{sheet}]: Testcase_Name and template_run_mode columns are required"
            )
        cases, names = [], set()
        for row_number, cells in enumerate(rows, 2):
            if all(_empty(cell.value) for cell in cells):
                continue
            try:
                if any(cell.data_type in ("f", "e") for cell in cells):
                    raise ValueError(
                        "formulas and Excel error cells are unsupported; paste values"
                    )
                if any(
                    not header and not _empty(cell.value)
                    for header, cell in zip(headers, cells)
                ):
                    raise ValueError("non-empty cell has no column header")
                params = _normalize(
                    {
                        header: cell.value
                        for header, cell in zip(headers, cells)
                        if header
                    }
                )
                name = params["Testcase_Name"]
                if name in names:
                    raise ValueError(f"duplicate Testcase_Name {name!r}")
                names.add(name)
                cases.append(params)
            except ValueError as exc:
                raise ValueError(f"{path} [{sheet}] row {row_number}: {exc}") from exc
        if not cases:
            raise ValueError(f"{path} [{sheet}]: no test cases")
        return cases
    finally:
        workbook.close()
